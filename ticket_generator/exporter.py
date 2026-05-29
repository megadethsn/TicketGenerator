import os
from collections import Counter
from datetime import datetime
from typing import Dict, Iterable, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    BLOCK_LABELS,
    BLOCK_PRACTICE,
    BLOCK_TEST,
    BLOCK_THEMATIC,
    CategoryGeneration,
    GenerationParams,
    Question,
    Ticket,
)


MONTHS_RU = {
    1: "января",
    2: "февраля",
    3: "марта",
    4: "апреля",
    5: "мая",
    6: "июня",
    7: "июля",
    8: "августа",
    9: "сентября",
    10: "октября",
    11: "ноября",
    12: "декабря",
}


HEADERS = [
    "Тема",
    "Название темы",
    "Вопрос",
    "Варианты ответов",
    "Правильный ответ",
    "Идентификатор",
]


def russian_date(date_value: datetime) -> str:
    return "%d %s %d" % (date_value.day, MONTHS_RU[date_value.month], date_value.year)


def category_filename(category: int, date_value: datetime) -> str:
    return "Билеты %d категория %s.xlsx" % (category, russian_date(date_value))


def stats_filename(date_value: datetime) -> str:
    return "Статистика генерации %s.xlsx" % russian_date(date_value)


def export_category_file(
    generation: CategoryGeneration, output_dir: str, date_value: datetime
) -> str:
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)

    for ticket in generation.tickets:
        worksheet = workbook.create_sheet("Билет %d" % ticket.number)
        _write_ticket_sheet(worksheet, ticket)

    identifiers = workbook.create_sheet("Идентификаторы")
    _write_identifiers_sheet(identifiers, generation.tickets)

    path = os.path.join(output_dir, category_filename(generation.category, date_value))
    workbook.save(path)
    return path


def export_stats_file(
    generations: Iterable[CategoryGeneration],
    output_dir: str,
    source_path: str,
    params: GenerationParams,
    date_value: datetime,
) -> str:
    generations = list(generations)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Сводка"
    summary.append(["Параметр", "Значение"])
    rows = [
        ("Исходный реестр", source_path),
        ("Дата генерации", date_value.strftime("%d.%m.%Y %H:%M")),
        ("Количество билетов", params.ticket_count),
        ("Основных билетов", params.main_ticket_count),
        ("Тестовых в билете", params.test_count),
        ("Тематических в билете", params.thematic_count),
        ("Практических в билете", params.practice_count),
    ]
    for row in rows:
        summary.append(list(row))
    _style_table(summary)

    usage_sheet = workbook.create_sheet("Использование")
    usage_sheet.append(["Категория", "Блок", "Идентификатор", "Использований"])
    for generation in generations:
        usage = Counter()
        block_by_id: Dict[str, str] = {}
        for ticket in generation.tickets:
            for block, questions in ticket.questions.items():
                for question in questions:
                    usage[question.question_id] += 1
                    block_by_id[question.question_id] = BLOCK_LABELS[block]
        for question_id in sorted(usage, key=int):
            usage_sheet.append(
                [
                    generation.category,
                    block_by_id[question_id],
                    question_id,
                    usage[question_id],
                ]
            )
    _style_table(usage_sheet)

    main_sheet = workbook.create_sheet("Основные билеты")
    main_sheet.append(["Категория", "Основные билеты"])
    for generation in generations:
        main = [str(ticket.number) for ticket in generation.tickets if ticket.is_main]
        main_sheet.append([generation.category, ", ".join(main)])
    _style_table(main_sheet)

    overlap_sheet = workbook.create_sheet("Пересечения")
    overlap_sheet.append(["Категория", "Билет A", "Билет B", "Общие тестовые", "Общие тематические", "Общие практические"])
    for generation in generations:
        tickets = generation.tickets
        for index, left in enumerate(tickets):
            for right in tickets[index + 1 :]:
                overlap_sheet.append(
                    [
                        generation.category,
                        left.number,
                        right.number,
                        len(set(left.ids(BLOCK_TEST)) & set(right.ids(BLOCK_TEST))),
                        len(set(left.ids(BLOCK_THEMATIC)) & set(right.ids(BLOCK_THEMATIC))),
                        len(set(left.ids(BLOCK_PRACTICE)) & set(right.ids(BLOCK_PRACTICE))),
                    ]
                )
    _style_table(overlap_sheet)

    warnings_sheet = workbook.create_sheet("Предупреждения")
    warnings_sheet.append(["Категория", "Предупреждение"])
    for generation in generations:
        for warning in generation.warnings:
            warnings_sheet.append([generation.category, warning])
    _style_table(warnings_sheet)

    path = os.path.join(output_dir, stats_filename(date_value))
    workbook.save(path)
    return path


def _write_ticket_sheet(worksheet, ticket: Ticket) -> None:
    worksheet.append(HEADERS)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for question in ticket.all_questions():
        worksheet.append(_question_row(question))
    _style_table(worksheet)


def _question_row(question: Question) -> List[str]:
    topic = question.topic
    if question.block == BLOCK_THEMATIC:
        topic = "Тематический вопрос"
    elif question.block == BLOCK_PRACTICE:
        topic = "Практическая задача"
    return [
        topic,
        question.topic_name,
        question.question,
        question.answers,
        question.correct_answer,
        question.question_id,
    ]


def _write_identifiers_sheet(worksheet, tickets: List[Ticket]) -> None:
    columns_per_block = 2
    rows_per_ticket = 6
    left_col = 1
    right_col = 3
    for index, ticket in enumerate(tickets):
        side_col = left_col if index < 5 else right_col
        block_index = index if index < 5 else index - 5
        row = block_index * rows_per_ticket + 1
        title = "Билет №%d" % ticket.number
        if ticket.is_main:
            title += " ★"
        worksheet.cell(row, side_col, title)
        worksheet.cell(row, side_col).font = Font(bold=True)
        worksheet.cell(row + 1, side_col, "Тестовые:")
        worksheet.cell(row + 1, side_col + 1, ", ".join(ticket.ids(BLOCK_TEST)))
        worksheet.cell(row + 2, side_col, "Тематические:")
        worksheet.cell(row + 2, side_col + 1, ", ".join(ticket.ids(BLOCK_THEMATIC)))
        worksheet.cell(row + 3, side_col, "Практические:")
        worksheet.cell(row + 3, side_col + 1, ", ".join(ticket.ids(BLOCK_PRACTICE)))
    _style_table(worksheet)
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 95
    worksheet.column_dimensions["C"].width = 18
    worksheet.column_dimensions["D"].width = 95


def _style_table(worksheet) -> None:
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {
        1: 18,
        2: 38,
        3: 55,
        4: 65,
        5: 50,
        6: 16,
    }
    for col_idx in range(1, worksheet.max_column + 1):
        width = widths.get(col_idx, 22)
        worksheet.column_dimensions[get_column_letter(col_idx)].width = width
    worksheet.freeze_panes = "A2"
