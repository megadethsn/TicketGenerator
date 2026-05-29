import re
from collections import Counter, defaultdict
from typing import Dict, Iterable, List, Tuple

from openpyxl import load_workbook

from .models import (
    BLOCK_PRACTICE,
    BLOCK_TEST,
    BLOCK_THEMATIC,
    Question,
    RegistryData,
    ValidationResult,
)


REQUIRED_HEADERS = {
    "Тема": "topic",
    "Название темы": "topic_name",
    "Вопрос": "question",
    "Варианты ответов": "answers",
    "Правильный ответ": "correct_answer",
    "Идентификатор": "question_id",
}


def normalize_header(value: object) -> str:
    return str(value or "").strip()


def normalize_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def category_sheet_name(category: int) -> str:
    return "Категория %d" % category


def block_from_id(question_id: str) -> str:
    clean_id = normalize_cell(question_id)
    if not re.fullmatch(r"\d+", clean_id) or len(clean_id) < 2:
        raise ValueError("идентификатор должен состоять из цифр и иметь минимум 2 знака")
    marker = clean_id[1]
    if marker == "1":
        return BLOCK_TEST
    if marker == "2":
        return BLOCK_THEMATIC
    if marker == "3":
        return BLOCK_PRACTICE
    raise ValueError("неизвестный блок по второй цифре идентификатора: %s" % marker)


def read_registry(path: str, categories: Iterable[int]) -> Tuple[RegistryData, ValidationResult]:
    errors: List[str] = []
    warnings: List[str] = []
    counts: Dict[int, Dict[str, int]] = defaultdict(lambda: defaultdict(int))
    questions_by_category: Dict[int, List[Question]] = {}

    try:
        workbook = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:
        return RegistryData({}, []), ValidationResult(
            ["Не удалось открыть файл Excel: %s" % exc], [], {}
        )

    for category in categories:
        sheet_name = category_sheet_name(category)
        if sheet_name not in workbook.sheetnames:
            errors.append("Категория %d: не найден лист '%s'." % (category, sheet_name))
            continue

        worksheet = workbook[sheet_name]
        header_map = _read_headers(worksheet)
        missing = [name for name in REQUIRED_HEADERS if name not in header_map]
        if missing:
            errors.append(
                "Категория %d: отсутствуют обязательные колонки: %s."
                % (category, ", ".join(missing))
            )
            continue

        category_questions: List[Question] = []
        ids: List[str] = []
        for row in range(2, worksheet.max_row + 1):
            if not _row_has_values(worksheet, row, header_map.values()):
                continue

            question_id = normalize_cell(
                worksheet.cell(row, header_map["Идентификатор"]).value
            )
            if not question_id:
                errors.append("%s!%d: пустой идентификатор." % (sheet_name, row))
                continue

            ids.append(question_id)
            if not question_id.startswith(str(category)):
                errors.append(
                    "%s!%d: идентификатор %s не соответствует категории %d."
                    % (sheet_name, row, question_id, category)
                )
                continue

            try:
                block = block_from_id(question_id)
            except ValueError as exc:
                errors.append("%s!%d: %s." % (sheet_name, row, exc))
                continue

            topic = normalize_cell(worksheet.cell(row, header_map["Тема"]).value)
            if block == BLOCK_THEMATIC:
                topic = "Тематический вопрос"
            elif block == BLOCK_PRACTICE:
                topic = "Практическая задача"

            question = Question(
                category=category,
                block=block,
                topic=topic,
                topic_name=normalize_cell(
                    worksheet.cell(row, header_map["Название темы"]).value
                ),
                question=normalize_cell(worksheet.cell(row, header_map["Вопрос"]).value),
                answers=normalize_cell(
                    worksheet.cell(row, header_map["Варианты ответов"]).value
                ),
                correct_answer=normalize_cell(
                    worksheet.cell(row, header_map["Правильный ответ"]).value
                ),
                question_id=question_id,
                source_sheet=sheet_name,
                source_row=row,
            )
            category_questions.append(question)
            counts[category][block] += 1

        duplicates = [qid for qid, amount in Counter(ids).items() if amount > 1]
        for question_id in duplicates:
            errors.append(
                "Категория %d: дублирующийся идентификатор %s." % (category, question_id)
            )

        questions_by_category[category] = category_questions

    return (
        RegistryData(questions_by_category=questions_by_category, warnings=warnings),
        ValidationResult(
            errors=errors,
            warnings=warnings,
            counts={cat: dict(value) for cat, value in counts.items()},
        ),
    )


def _read_headers(worksheet) -> Dict[str, int]:
    result: Dict[str, int] = {}
    for col in range(1, worksheet.max_column + 1):
        header = normalize_header(worksheet.cell(1, col).value)
        if header in REQUIRED_HEADERS:
            result[header] = col
    return result


def _row_has_values(worksheet, row: int, columns: Iterable[int]) -> bool:
    for col in columns:
        value = worksheet.cell(row, col).value
        if value not in (None, ""):
            return True
    return False
